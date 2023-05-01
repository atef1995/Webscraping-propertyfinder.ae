import csv
from docx import Document

# Open the CSV file and read the data
with open('output.csv', 'r') as csvfile:
    reader = csv.reader(csvfile)
    data = list(reader)

# Create a new Word document
document = Document()

# Add a table to the document with the data from the CSV file
table = document.add_table(rows=len(data), cols=len(data[0]))

# Populate the table with the data from the CSV file
for i, row in enumerate(data):
    for j, cell in enumerate(row):
        table.cell(i, j).text = cell

# Save the document
document.save('mydocument.docx')
